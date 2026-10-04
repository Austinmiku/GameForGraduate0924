"""Export the game's question bank from the supplied Excel workbook."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


HEADERS = [
    "關卡編號", "關卡名稱", "知識單元", "題號", "題目",
    "選項 A", "選項 B", "選項 C", "選項 D", "題庫正確選項",
    "正確答案", "提示", "答案解析",
]


def export(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, data_only=False, keep_vba=False)
    sheet = workbook["完整題庫"]
    if [sheet.cell(5, column).value for column in range(1, 14)] != HEADERS:
        raise ValueError("題庫欄位與預期格式不符")

    bank = {str(level): [] for level in range(1, 9)}
    for row_number in range(6, 86):
        cells = [sheet.cell(row_number, column) for column in range(1, 14)]
        if any(cell.data_type == "f" for cell in cells):
            raise ValueError(f"第 {row_number} 列含公式，請先轉成固定文字")
        level, title, unit, number, question, *rest = [cell.value for cell in cells]
        a, b, c, d, correct_letter, correct_text, hint, explanation = rest
        if not isinstance(level, (int, float)) or int(level) not in range(1, 9):
            raise ValueError(f"第 {row_number} 列的關卡編號無效")
        level = int(level)
        if number != len(bank[str(level)]) + 1:
            raise ValueError(f"第 {row_number} 列的題號或順序無效")
        if any(not isinstance(value, str) or not value for value in
               (title, unit, question, a, b, c, d, correct_letter, correct_text, hint, explanation)):
            raise ValueError(f"第 {row_number} 列有空白或無效的題目欄位")
        if correct_letter not in "ABCD":
            raise ValueError(f"第 {row_number} 列的正確選項無效")
        answer_index = "ABCD".index(correct_letter)
        options = [a, b, c, d]
        if options[answer_index] != correct_text:
            raise ValueError(f"第 {row_number} 列的正確答案與選項不一致")
        bank[str(level)].append([question, options, answer_index, hint, explanation])

    if any(len(questions) != 10 for questions in bank.values()):
        raise ValueError("每關必須恰好有 10 題")
    for row_number in range(86, sheet.max_row + 1):
        if any(sheet.cell(row_number, column).value is not None for column in range(1, 14)):
            raise ValueError(f"第 {row_number} 列在預期題庫範圍外，請先確認題數")
    for level, questions in bank.items():
        if len({question[0] for question in questions}) != 10:
            raise ValueError(f"第 {level} 關有重複題目")

    content = (
        "// 由 Excel 題庫產生；請修改 Excel 後重新匯出此檔。\n"
        "window.GUARDIAN_QUESTION_BANK = "
        + json.dumps(bank, ensure_ascii=False, indent=2)
        + ";\n"
    )
    destination.write_text(content, encoding="utf-8")
    print(f"已匯出 {sum(map(len, bank.values()))} 題至 {destination}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Excel .xlsm 題庫檔")
    parser.add_argument(
        "--output", type=Path,
        default=Path(__file__).resolve().parent.parent / "question-bank.js",
        help="輸出的 JS 題庫檔",
    )
    args = parser.parse_args()
    export(args.source, args.output)
